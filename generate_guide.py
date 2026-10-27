import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_document():
    doc = Document()

    # Set page margins (Normal: 1 inch / 2.54 cm)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    def set_cell_background(cell, hex_color):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
        tcPr.append(shd)

    def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f'''
            <w:tcMar {nsdecls("w")}>
                <w:top w:w="{top}" w:type="dxa"/>
                <w:bottom w:w="{bottom}" w:type="dxa"/>
                <w:left w:w="{left}" w:type="dxa"/>
                <w:right w:w="{right}" w:type="dxa"/>
            </w:tcMar>
        ''')
        tcPr.append(tcMar)

    def format_run(run, font_name="Khmer OS Battambang", size_pt=11, bold=False, italic=False, color_rgb=(30, 41, 59)):
        run.font.name = font_name
        run.font.size = Pt(size_pt)
        run.bold = bold
        run.italic = italic
        run.font.color.rgb = RGBColor(*color_rgb)
        rPr = run._r.get_or_add_rPr()
        rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="{font_name}" w:hAnsi="{font_name}" w:cs="{font_name}"/>')
        rPr.append(rFonts)

    def add_callout(title, text, bg_color="F0F9FF", border_color="0284C7"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        tbl.columns[0].width = Inches(6.5)
        
        cell = tbl.cell(0, 0)
        set_cell_background(cell, bg_color)
        set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
        
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:left w:val="single" w:sz="36" w:space="0" w:color="{border_color}"/>
                <w:top w:val="none"/>
                <w:right w:val="none"/>
                <w:bottom w:val="none"/>
            </w:tcBorders>
        ''')
        tcPr.append(borders)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.25
        run_title = p.add_run(f"📌 {title}: ")
        format_run(run_title, font_name="Khmer OS Battambang", size_pt=11, bold=True, color_rgb=(2, 132, 199))
        
        run_text = p.add_run(text)
        format_run(run_text, font_name="Khmer OS Battambang", size_pt=10.5, color_rgb=(51, 65, 85))
        
        p_after = doc.add_paragraph()
        p_after.paragraph_format.space_before = Pt(0)
        p_after.paragraph_format.space_after = Pt(4)

    def add_heading_1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(20)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        format_run(run, font_name="Khmer OS Muol Light", size_pt=13.5, bold=True, color_rgb=(15, 23, 42))
        
        pPr = p._p.get_or_add_pPr()
        pBdr = parse_xml(f'''
            <w:pBdr {nsdecls("w")}>
                <w:bottom w:val="single" w:sz="16" w:space="4" w:color="0284C7"/>
            </w:pBdr>
        ''')
        pPr.append(pBdr)

    def add_heading_2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        format_run(run, font_name="Khmer OS Battambang", size_pt=12, bold=True, color_rgb=(2, 132, 199))

    def add_heading_3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(9)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        format_run(run, font_name="Khmer OS Battambang", size_pt=11, bold=True, color_rgb=(71, 85, 105))

    def add_body(text, bold_prefix="", space_after=4):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.25
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            format_run(r_pre, font_name="Khmer OS Battambang", size_pt=11, bold=True, color_rgb=(30, 41, 59))
        r_body = p.add_run(text)
        format_run(r_body, font_name="Khmer OS Battambang", size_pt=11, color_rgb=(51, 65, 85))

    def add_bullet(text, bold_prefix=""):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.25
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            format_run(r_pre, font_name="Khmer OS Battambang", size_pt=10.5, bold=True, color_rgb=(15, 23, 42))
        r_body = p.add_run(text)
        format_run(r_body, font_name="Khmer OS Battambang", size_pt=10.5, color_rgb=(51, 65, 85))

    # ================= COVER / TITLE SECTION =================
    p_top = doc.add_paragraph()
    p_top.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_top.paragraph_format.space_after = Pt(2)
    r1 = p_top.add_run("ព្រះរាជាណាចក្រកម្ពុជា\nជាតិ សាសនា ព្រះមហាក្សត្រ\n***")
    format_run(r1, font_name="Khmer OS Muol Light", size_pt=13, bold=True, color_rgb=(15, 23, 42))

    p_sch = doc.add_paragraph()
    p_sch.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sch.paragraph_format.space_before = Pt(10)
    p_sch.paragraph_format.space_after = Pt(14)
    r2 = p_sch.add_run("សាកលវិទ្យាល័យជាតិបាត់ដំបង (NUBB)\nប្រព័ន្ធគ្រប់គ្រងថ្នាក់រៀនអនឡាញ (Class Management System 2027)")
    format_run(r2, font_name="Khmer OS Muol Light", size_pt=13, bold=True, color_rgb=(2, 132, 199))

    # Big Banner Box
    tbl_cover = doc.add_table(rows=1, cols=1)
    tbl_cover.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_cover.autofit = False
    tbl_cover.columns[0].width = Inches(6.5)
    c_cover = tbl_cover.cell(0, 0)
    set_cell_background(c_cover, "0F172A")
    set_cell_margins(c_cover, top=260, bottom=260, left=250, right=250)

    p_cover = c_cover.paragraphs[0]
    p_cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cover.paragraph_format.line_spacing = 1.3
    r_cov_title = p_cover.add_run("សៀវភៅមគ្គទេសក៍ និងមេរៀនណែនាំលម្អិត\nស្តីពីការចូលប្រើប្រាស់ និងគ្រប់គ្រងប្រព័ន្ធសាថ្មី\n")
    format_run(r_cov_title, font_name="Khmer OS Muol Light", size_pt=16, bold=True, color_rgb=(255, 255, 255))
    r_cov_sub = p_cover.add_run("ការកែលម្អប្រព័ន្ធគ្រប់គ្រងថ្នាក់រៀនជំនាន់ចុងក្រោយ សម្រាប់លោកគ្រូ-អ្នកគ្រូបង្រៀន និងគណៈគ្រប់គ្រង")
    format_run(r_cov_sub, font_name="Khmer OS Battambang", size_pt=12, bold=False, color_rgb=(56, 189, 248))

    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_before = Pt(14)
    p_meta.paragraph_format.space_after = Pt(18)
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_meta = p_meta.add_run("🌐 អាសយដ្ឋានផ្លូវការ៖ https://class.mitpukae.com\n"
                           "✍️ រៀបចំ និងចងក្រងដោយ៖ លោកគ្រូ សំ គឹមឆេង (Sam Kimchheng)\n"
                           "✈️ Telegram ជំនួយបច្ចេកទេសផ្ទាល់៖ https://t.me/Kimchheng12 (ID: @Kimchheng12)\n"
                           "📅 ឆ្នាំសិក្សា៖ ២០២៦ - ២០២៧ (ធ្វើបច្ចុប្បន្នភាពចុងក្រោយ)")
    format_run(r_meta, font_name="Khmer OS Battambang", size_pt=10.5, color_rgb=(71, 85, 105))

    doc.add_page_break()

    # ================= TABLE OF CONTENTS =================
    add_heading_1("មាតិកានៃសៀវភៅណែនាំសាថ្មី (TABLE OF CONTENTS)")
    
    toc_items = [
        ("ជំពូកទី ១", "សេចក្តីផ្តើម និងព័ត៌មានទូទៅអំពីប្រព័ន្ធ"),
        ("ជំពូកទី ២", "របៀបចូលប្រើប្រាស់ប្រព័ន្ធ (Login) និងការគ្រប់គ្រងគណនី"),
        ("ជំពូកទី ៣", "ការគ្រប់គ្រងឆ្នាំសិក្សា (Multi-Academic Year Management)"),
        ("ជំពូកទី ៤", "ការបំពេញព័ត៌មានផ្ទាល់ខ្លួនគ្រូ ហត្ថលេខា និងត្រាសាលា (Teacher Profile & Signatures)"),
        ("ជំពូកទី ៥", "ការគ្រប់គ្រងបញ្ជីថ្នាក់រៀនផ្ទាល់ខ្លួន (Classroom Management)"),
        ("ជំពូកទី ៦", "ការគ្រប់គ្រងបញ្ជីឈ្មោះសិស្ស (Students Management)"),
        ("ជំពូកទី ៧", "ការកត់ត្រា និងតាមដានអវត្តមានសិស្ស (Daily & Monthly Attendance)"),
        ("ជំពូកទី ៨", "ការបញ្ចូល និងគ្រប់គ្រងពិន្ទុសិស្ស (Scores & Grading Rules)"),
        ("ជំពូកទី ៩", "ស្ថិតិ និងការវិភាគទិន្នន័យសិក្សា (Student Statistics & Subject Breakdown)"),
        ("ជំពូកទី ១០", "វិញ្ញាបនបត្របញ្ជាក់ការសិក្សា និងការផ្ទៀងផ្ទាត់ QR Code (Certificates & Verification)"),
        ("ជំពូកទី ១១", "ការគ្រប់គ្រងវិញ្ញាសា និងការប្រឡងអនឡាញ (Online Exams)"),
        ("ជំពូកទី ១២", "ការគ្រប់គ្រងមេរៀន និងកាលវិភាគបង្រៀន (Lessons & Schedules)"),
        ("ជំពូកទី ១៣", "ការទាញយក និងបោះពុម្ពរបាយការណ៍សរុប (Reports & Printing)"),
        ("ជំពូកទី ១៤", "ការគ្រប់គ្រងគណនីគ្រូបង្រៀនសម្រាប់ Admin (Admin Accounts Management)"),
        ("ជំពូកទី ១៥", "ផ្នែកទំនាក់ទំនង ជំនួយបច្ចេកទេស និងសុវត្ថិភាព (Support & Security)")
    ]

    tbl_toc = doc.add_table(rows=1, cols=2)
    tbl_toc.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_toc.autofit = False
    tbl_toc.columns[0].width = Inches(1.8)
    tbl_toc.columns[1].width = Inches(4.7)
    
    # Header row
    hdr_cells = tbl_toc.rows[0].cells
    hdr_cells[0].text = "ជំពូក / លេខរៀង"
    hdr_cells[1].text = "ខ្លឹមសារមេរៀនណែនាំ"
    for i, c in enumerate(hdr_cells):
        set_cell_background(c, "0284C7")
        set_cell_margins(c, top=100, bottom=100, left=120, right=120)
        p = c.paragraphs[0]
        for r in p.runs:
            format_run(r, font_name="Khmer OS Battambang", size_pt=11, bold=True, color_rgb=(255, 255, 255))
    
    for ch, desc in toc_items:
        row_cells = tbl_toc.add_row().cells
        row_cells[0].text = ch
        row_cells[1].text = desc
        for i, c in enumerate(row_cells):
            set_cell_background(c, "F8FAFC" if toc_items.index((ch, desc)) % 2 == 0 else "FFFFFF")
            set_cell_margins(c, top=80, bottom=80, left=120, right=120)
            p = c.paragraphs[0]
            for r in p.runs:
                format_run(r, font_name="Khmer OS Battambang", size_pt=10.5, bold=(i == 0), color_rgb=(30, 41, 59))

    # ================= CHAPTER 1 =================
    add_heading_1("ជំពូកទី ១៖ សេចក្តីផ្តើម និងព័ត៌មានទូទៅអំពីប្រព័ន្ធ")
    add_body("ប្រព័ន្ធគ្រប់គ្រងថ្នាក់រៀនអនឡាញ (Class Management System 2027) ត្រូវបានអភិវឌ្ឍ និងកែលម្អសាថ្មីយ៉ាងទូលំទូលាយ ដើម្បីផ្តល់ជូនលោកគ្រូ-អ្នកគ្រូវិធីសាស្ត្រគ្រប់គ្រងថ្នាក់រៀនដ៏ទំនើប រហ័ស និងមានប្រសិទ្ធភាពខ្ពស់បំផុត។ កំណែថ្មីនេះបានពង្រឹងនូវប្រព័ន្ធសុវត្ថិភាពទិន្នន័យឯកោដាច់ដោយឡែករវាងគ្រូម្នាក់ៗ ការគ្រប់គ្រងឆ្នាំសិក្សាចម្រុះ ការស្រង់អវត្តមាន និងពិន្ទុច្បាស់លាស់ ព្រមទាំងការចេញវិញ្ញាបនបត្រស្វ័យប្រវត្តិជាមួយប្រព័ន្ធ QR Code ផ្ទៀងផ្ទាត់សច្ចភាព។")
    
    add_heading_2("១.១. អាសយដ្ឋានតំណភ្ជាប់ផ្លូវការ (Official URLs)")
    add_body("លោកគ្រូ-អ្នកគ្រូ និងសិស្សានុសិស្សអាចចូលប្រើប្រាស់ប្រព័ន្ធពីគ្រប់ទីកន្លែង និងគ្រប់ពេលវេលា តាមរយៈអាសយដ្ឋានផ្លូវការ៖")
    add_bullet("https://class.mitpukae.com (ទំព័រទូទៅ និងការចូលរួម)", "👉 គេហទំព័រផ្លូវការ៖ ")
    add_bullet("https://class.mitpukae.com/login (សម្រាប់លោកគ្រូ-អ្នកគ្រូ និងគណៈគ្រប់គ្រង Login)", "👉 Link Login គ្រូ៖ ")
    add_bullet("https://class.mitpukae.com/test (សម្រាប់សិស្សចូលធ្វើតេស្តប្រឡងអនឡាញ)", "👉 Link សិស្សប្រឡង៖ ")
    add_bullet("https://class.mitpukae.com/certificate/verify/{serial} (សម្រាប់ស្កេនផ្ទៀងផ្ទាត់វិញ្ញាបនបត្រ)", "👉 Link ផ្ទៀងផ្ទាត់វិញ្ញាបនបត្រ៖ ")
    
    add_heading_2("១.២. តម្រូវការបច្ចេកទេស និងឧបករណ៍ប្រើប្រាស់")
    add_bullet("ដំណើរការបានយ៉ាងរលូនលើគ្រប់ឧបករណ៍៖ កុំព្យូទ័រលើតុ (Desktop PC), កុំព្យូទ័រយួរដៃ (Laptop), ថេប្លេត (iPad/Tablet), និងទូរស័ព្ទដៃឆ្លាតវៃ (iOS & Android)។")
    add_bullet("កម្មវិធីរុករក (Web Browser) ដែលណែនាំ៖ Google Chrome, Microsoft Edge, Safari, ឬ Mozilla Firefox (កំណែទម្រង់ថ្មី)។")
    add_bullet("តម្រូវឱ្យមានការតភ្ជាប់បណ្តាញអ៊ីនធឺណិត (Wi-Fi ឬ 4G/5G) ដើម្បីរក្សាទុក និងធ្វើបច្ចុប្បន្នភាពទិន្នន័យលើពពក (Cloud Server) ភ្លាមៗ។")

    add_callout("ចំណុចពិសេស និងការការពារទិន្នន័យ", "ប្រព័ន្ធជំនាន់ថ្មី ២០២៧ ត្រូវបានរចនាឡើងជាមួយនឹងស្ថាបត្យកម្ម Multi-Tenant Isolation ១០០% ដែលធានាថាគ្រូនីមួយៗមានទិន្នន័យផ្ទាល់ខ្លួនដាច់ដោយឡែក មិនប៉ះពាល់ មិនឃើញ និងមិនជាន់គ្នាជាមួយគ្រូដទៃឡើយ ទោះបីជាបង្រៀនមុខវិជ្ជាដូចគ្នា ឬឈ្មោះថ្នាក់ដូចគ្នាក៏ដោយ។")

    # ================= CHAPTER 2 =================
    add_heading_1("ជំពូកទី ២៖ របៀបចូលប្រើប្រាស់ប្រព័ន្ធ (Login) និងការគ្រប់គ្រងគណនី")
    add_body("ដើម្បីធានាសុវត្ថិភាពព័ត៌មានផ្ទាល់ខ្លួន ពិន្ទុ និងវត្តមានរបស់សិស្ស មានតែលោកគ្រូ-អ្នកគ្រូដែលមានគណនីផ្លូវការប៉ុណ្ណោះ ទើបអាចចូលប្រើប្រាស់មុខងារបង្រៀនបាន។")

    add_heading_2("២.១. ជំហានក្នុងការ Login ចូលប្រព័ន្ធ")
    add_bullet("បើកកម្មវិធី Browser រួចវាយអាសយដ្ឋាន៖ https://class.mitpukae.com/login", "ជំហានទី ១៖ ")
    add_bullet("ផ្ទាំង «ចូលប្រើប្រព័ន្ធគ្រប់គ្រង» នឹងបង្ហាញនៅចំកណ្តាលអេក្រង់។", "ជំហានទី ២៖ ")
    add_bullet("បញ្ចូល ឈ្មោះអ្នកប្រើប្រាស់ (Username) របស់អ្នក (ឧទាហរណ៍៖ samkimchheng ឬ seungsomuny)។", "ជំហានទី ៣៖ ")
    add_bullet("បញ្ចូល លេខសម្ងាត់ (Password) របស់អ្នក។", "ជំហានទី ៤៖ ")
    add_bullet("អាចចុចលើរូបសញ្ញា ភ្នែក 👁️ នៅខាងស្តាំប្រអប់ដើម្បីផ្ទៀងផ្ទាត់លេខសម្ងាត់មុននឹងចុចចូល។", "ជំហានទី ៥៖ ")
    add_bullet("ចុចលើប៊ូតុងពណ៌ខៀវ «ចូលប្រព័ន្ធ (Login)» ដើម្បីបញ្ចប់។", "ជំហានទី ៦៖ ")

    add_heading_2("២.២. ការបែងចែកសិទ្ធិប្រើប្រាស់រវាង Admin និង Teacher")
    add_bullet("គណនីអ្នកគ្រប់គ្រង (Admin)៖ មានសិទ្ធិគ្រប់គ្រងគណនីគ្រូទាំងអស់ បង្កើតគ្រូថ្មី បើក/ផ្អាកគណនី និងមើលស្ថិតិរួមនៃប្រព័ន្ធ។", "១. ROLE_ADMIN៖ ")
    add_bullet("គណនីគ្រូបង្រៀន (Teacher)៖ គ្រប់គ្រងថ្នាក់រៀនផ្ទាល់ខ្លួន បញ្ជីសិស្ស វត្តមាន ពិន្ទុ វិញ្ញាសា កាលវិភាគ និងវិញ្ញាបនបត្ររបស់ខ្លួន។", "២. ROLE_TEACHER៖ ")

    add_heading_2("២.៣. មុខងារចងចាំការចូលប្រើប្រាស់ (Persistent Session)")
    add_body("ប្រព័ន្ធបានបំពាក់មុខងារចងចាំគណនីស្វ័យប្រវត្តិតាម Cookie រយៈពេល ៣០ ថ្ងៃ។ នៅពេលលោកគ្រូ-អ្នកគ្រូបាន Login ម្តងរួចហើយ លើកក្រោយបើក Browser ឡើងវិញ អាចចូលធ្វើការបានភ្លាមៗដោយមិនបាច់វាយលេខសម្ងាត់ម្តងទៀតឡើយ។")

    # ================= CHAPTER 3 =================
    add_heading_1("ជំពូកទី ៣៖ ការគ្រប់គ្រងឆ្នាំសិក្សា (Multi-Academic Year Management)")
    add_body("ប្រព័ន្ធគ្រប់គ្រងថ្នាក់រៀន ២០២៧ គាំទ្រការគ្រប់គ្រងទិន្នន័យច្រើនឆ្នាំសិក្សា (Multi-Academic Year) ដោយរលូន និងមានសុវត្ថិភាពខ្ពស់។ មុខងារនេះជួយឱ្យលោកគ្រូ-អ្នកគ្រូអាចរក្សាទុកទិន្នន័យចាស់ៗបានគង់វង្ស និងចាប់ផ្តើមឆ្នាំសិក្សាថ្មីដោយមិនបាត់បង់ទិន្នន័យឆ្នាំមុនឡើយ។")

    add_heading_2("៣.១. របារជ្រើសរើសឆ្នាំសិក្សា (Top Bar Academic Year Selector)")
    add_body("នៅផ្នែកខាងលើនៃអេក្រង់ (Top Navigation Bar) លើគ្រប់ទំព័រទាំងអស់ មានប្រអប់ជ្រើសរើសឆ្នាំសិក្សា៖")
    add_bullet("បង្ហាញឆ្នាំសិក្សាបច្ចុប្បន្ន (ឧទាហរណ៍៖ ២០២៦-២០២៧)។")
    add_bullet("លោកគ្រូ-អ្នកគ្រូអាចចុចលើប្រអប់ទម្លាក់ចុះ (Dropdown) ដើម្បីជ្រើសរើសឆ្នាំសិក្សាផ្សេងទៀត (ឧទាហរណ៍៖ ២០២៥-២០២៦) ដើម្បីពិនិត្យឡើងវិញនូវថ្នាក់រៀន សិស្ស វត្តមាន និងពិន្ទុចាស់ៗ។")
    add_bullet("នៅពេលលោកគ្រូ-អ្នកគ្រូជ្រើសរើសឆ្នាំណាមួយ ប្រព័ន្ធនឹងផ្លាស់ប្តូរទិន្នន័យទាំងស្រុងទៅតាមឆ្នាំសិក្សានោះដោយស្វ័យប្រវត្តិ។")

    add_heading_2("៣.២. ការបង្កើតឆ្នាំសិក្សាថ្មី (+ បង្កើតឆ្នាំសិក្សាថ្មី)")
    add_body("នៅពេលឈានចូលឆ្នាំសិក្សាថ្មី លោកគ្រូ-អ្នកគ្រូអាចបង្កើតឆ្នាំសិក្សាថ្មីបានយ៉ាងងាយស្រួល៖")
    add_bullet("ចូលទៅកាន់ទំព័រ «បញ្ជីថ្នាក់រៀនខ្ញុំ» ឬចុចលើប៊ូតុង «+ បង្កើតឆ្នាំសិក្សាថ្មី»។", "ជំហានទី ១៖ ")
    add_bullet("វាយបញ្ចូលឈ្មោះឆ្នាំសិក្សាថ្មី (ឧទាហរណ៍៖ ២០២៧-២០២៨)។", "ជំហានទី ២៖ ")
    add_bullet("ចុច «រក្សាទុក»។ ឆ្នាំសិក្សាថ្មីនឹងត្រូវបានបង្កើតឡើងភ្លាមៗ។", "ជំហានទី ៣៖ ")

    add_heading_2("៣.៣. មុខងារថតចម្លងថ្នាក់រៀនពីឆ្នាំមុន (Copy Classrooms to New Year)")
    add_body("ដើម្បីកុំឱ្យលោកគ្រូ-អ្នកគ្រូចំណាយពេលវាយបញ្ចូលឈ្មោះថ្នាក់រៀនសាឡើងវិញនៅឆ្នាំថ្មី ប្រព័ន្ធមានប៊ូតុង «ថតចម្លងថ្នាក់ពីឆ្នាំមុន»៖")
    add_bullet("ប្រព័ន្ធនឹងទាញយកឈ្មោះថ្នាក់រៀនទាំងអស់ពីឆ្នាំសិក្សាចាស់ មកបង្កើតក្នុងឆ្នាំសិក្សាថ្មីជូនលោកគ្រូ-អ្នកគ្រូដោយស្វ័យប្រវត្តិ។")
    add_bullet("បញ្ជីឈ្មោះសិស្ស វត្តមាន និងពិន្ទុនៃឆ្នាំថ្មី នឹងចាប់ផ្តើមជាទំព័រស្អាត (Clean Slate) ដើម្បីឱ្យគ្រូបន្ថែមសិស្សថ្មីចូលរៀន។")

    # ================= CHAPTER 4 =================
    add_heading_1("ជំពូកទី ៤៖ ការបំពេញព័ត៌មានផ្ទាល់ខ្លួនគ្រូ ហត្ថលេខា និងត្រាសាលា")
    add_body("ទំព័រ «បំពេញព័ត៌មានផ្ទាល់ខ្លួន (/profile)» គឺជាចំណុចស្នូលក្នុងការកំណត់អត្តសញ្ញាណរបស់លោកគ្រូ-អ្នកគ្រូ និងស្ថាប័នអប់រំ។ ព័ត៌មានដែលបានបំពេញនៅទីនេះ នឹងត្រូវយកទៅប្រើប្រាស់ស្វ័យប្រវត្តិនៅលើក្បាលរបាយការណ៍ វិញ្ញាបនបត្រ និងទំនាក់ទំនង។")

    add_heading_2("៤.១. ព័ត៌មានទូទៅរបស់គ្រូបង្រៀន")
    add_bullet("គោត្តនាម និងនាម (FullName)៖ បញ្ចូលឈ្មោះពេញរបស់គ្រូជាភាសាខ្មែរ (ឧទាហរណ៍៖ សំ គឹមឆេង)។")
    add_bullet("ភេទ (Gender)៖ ជ្រើសរើស ប្រុស ឬ ស្រី។")
    add_bullet("ឈ្មោះគ្រឹះស្ថានសិក្សា / សាលារៀន (School Name)៖ បញ្ចូលឈ្មោះសាលា (ឧទាហរណ៍៖ វិទ្យាល័យ នេត យ៉ង់ ឬ សាកលវិទ្យាល័យជាតិបាត់ដំបង)។")
    add_bullet("រាជធានី / ខេត្ត (Province)៖ បញ្ចូលទីតាំងខេត្ត (ឧទាហរណ៍៖ ខេត្តបាត់ដំបង)។")
    add_bullet("មុខវិជ្ជាបង្រៀន (Teaching Subject)៖ ឧទាហរណ៍៖ រូបវិទ្យា, គណិតវិទ្យា, ព័ត៌មានវិទ្យា។")
    add_bullet("ថ្នាក់បង្រៀន (Teaching Classes)៖ បញ្ជាក់ថ្នាក់ដែលទទួលបន្ទុក (ឧទាហរណ៍៖ ថ្នាក់ទី១០ A, ១១ B)។")
    add_bullet("លេខទូរស័ព្ទ (Phone Number)៖ សម្រាប់ទំនាក់ទំនងការងារ។")

    add_heading_2("៤.២. ការភ្ជាប់គណនីតេឡេក្រាមផ្ទាល់ខ្លួន (Personal Telegram Integration)")
    add_body("គ្រូនីមួយៗអាចបញ្ចូល គណនីតេឡេក្រាមផ្ទាល់ខ្លួន របស់គាត់បាន (ឧទាហរណ៍៖ Kimchheng12 ឬ somuny_teacher)៖")
    add_bullet("នៅពេលគ្រូបានបញ្ចូលគណនី Telegram របស់គាត់ហើយ ប្រព័ន្ធនឹងធ្វើបច្ចុប្បន្នភាពតំណភ្ជាប់ Telegram នៅលើ Sidebar ខាងឆ្វេង និងផ្ទាំង Dashboard របស់គ្រូនោះដោយស្វ័យប្រវត្តិ។")
    add_bullet("គ្រូផ្សេងៗទៀតនឹងមិនជាប់តេឡេក្រាមរបស់គ្រូដទៃឡើយ ដោយគ្រូនីមួយៗមានគណនីតេឡេក្រាមរៀងៗខ្លួន។")

    add_heading_2("៤.៣. ការបញ្ចូលហត្ថលេខាគ្រូ ហត្ថលេខានាយក និងត្រាសាលា")
    add_body("ប្រព័ន្ធបានរៀបចំផ្ទាំងគូរ និង Upload ឌីជីថលយ៉ាងទំនើប៖")
    add_bullet("ហត្ថលេខាគ្រូបង្រៀន (Teacher Signature)៖ អាចគូសដោយផ្ទាល់ដៃលើអេក្រង់ (Signature Drawing Pad) ឬចុច Upload ជារូបភាពហត្ថលេខាដែលមានស្រាប់។")
    add_bullet("ហត្ថលេខានាយកសាលា (Principal Signature)៖ អាចគូស ឬ Upload រូបភាពហត្ថលេខានាយកសាលា។")
    add_bullet("ត្រាសាលារៀន (Official School Stamp)៖ Upload រូបភាពត្រាមូល ឬត្រាផ្លូវការរបស់សាលារៀន (ទម្រង់ PNG ឬ JPG)។")
    add_bullet("បន្ទាប់ពីបំពេញរួចរាល់ ចុចប៊ូតុង «រក្សាទុកព័ត៌មានផ្ទាល់ខ្លួន»។ ប្រព័ន្ធនឹងជូនដំណឹងជោគជ័យភ្លាមៗ។")

    add_callout("អត្ថប្រយោជន៍នៃត្រា និងហត្ថលេខាឌីជីថល", "រូបភាពត្រាសាលា និងហត្ថលេខាដែលលោកគ្រូ-អ្នកគ្រូបានរក្សាទុកក្នុងទំព័រព័ត៌មានផ្ទាល់ខ្លួន នឹងត្រូវបានតភ្ជាប់ដោយស្វ័យប្រវត្តិចូលទៅកាន់ «វិញ្ញាបនបត្របញ្ជាក់ការសិក្សា» របស់សិស្ស ដោយមិនចាំបាច់បោះត្រា ឬស៊ីញ៉េដោយដៃម្តងមួយៗលើក្រដាសឡើយ!")

    # ================= CHAPTER 5 =================
    add_heading_1("ជំពូកទី ៥៖ ការគ្រប់គ្រងបញ្ជីថ្នាក់រៀនផ្ទាល់ខ្លួន (Classroom Management)")
    add_body("លោកគ្រូ-អ្នកគ្រូអាចបង្រៀនច្រើនថ្នាក់ក្នុងពេលតែមួយ។ ប្រព័ន្ធគ្រប់គ្រងថ្នាក់រៀន ២០២៧ ធានាថាគ្រូនីមួយៗមានបញ្ជីថ្នាក់រៀនផ្ទាល់ខ្លួនដាច់ដោយឡែក ១០០%។")

    add_heading_2("៥.១. ការបន្ថែមថ្នាក់រៀនថ្មី")
    add_bullet("ចូលទៅកាន់ម៉ឺនុយ «បញ្ជីថ្នាក់រៀនខ្ញុំ (Classrooms)»។", "ជំហានទី ១៖ ")
    add_bullet("ចុចលើប៊ូតុងពណ៌ខៀវ «+ បន្ថែមថ្នាក់រៀនថ្មី»។", "ជំហានទី ២៖ ")
    add_bullet("វាយបញ្ចូល ឈ្មោះថ្នាក់រៀន (ឧទាហរណ៍៖ ថ្នាក់ទី១០ A, ថ្នាក់ទី១១ វិទ្យាសាស្ត្រ, ឬ IT-Year3)។", "ជំហានទី ៣៖ ")
    add_bullet("វាយបញ្ចូល បន្ទប់សិក្សា ឬវេនសិក្សា (បើមាន)។", "ជំហានទី ៤៖ ")
    add_bullet("ចុច «រក្សាទុក (Save)»។ ថ្នាក់រៀននឹងបង្ហាញក្នុងបញ្ជីភ្លាមៗ។", "ជំហានទី ៥៖ ")

    add_heading_2("៥.២. ការកែសម្រួល ឬលុបថ្នាក់រៀន")
    add_bullet("ការកែប្រែ (Edit)៖ ចុចលើប៊ូតុង «កែប្រែ ✏️» ដើម្បីផ្លាស់ប្តូរឈ្មោះថ្នាក់ ឬបន្ទប់សិក្សា។")
    add_bullet("ការលុប (Delete)៖ ចុចលើប៊ូតុង «លុប 🗑️» ប្រសិនបើលែងបង្រៀនថ្នាក់នោះ។ (ប្រព័ន្ធនឹងស្នើសុំការបញ្ជាក់មុននឹងលុប ដើម្បីជៀសវាងការច្រឡំដៃ)។")

    # ================= CHAPTER 6 =================
    add_heading_1("ជំពូកទី ៦៖ ការគ្រប់គ្រងបញ្ជីឈ្មោះសិស្ស (Students Management)")
    add_body("ទិន្នន័យសិស្សគឺជាមូលដ្ឋានគ្រឹះសម្រាប់ដំណើរការស្រង់វត្តមាន បញ្ចូលពិន្ទុ និងការចេញវិញ្ញាបនបត្រ។ ក្នុងប្រព័ន្ធជំនាន់ថ្មី ទិន្នន័យគំរូចាស់ៗត្រូវបានលុបចេញស្អាត ដើម្បីឱ្យលោកគ្រូ-អ្នកគ្រូបញ្ចូលសិស្សពិតប្រាកដរបស់ខ្លួន។")

    add_heading_2("៦.១. របៀបបន្ថែមសិស្សថ្មី")
    add_bullet("ចុចលើម៉ឺនុយ «បញ្ជីឈ្មោះសិស្ស (Students)» -> រួចចុចប៊ូតុង «+ បន្ថែមសិស្សថ្មី»។", "ជំហានទី ១៖ ")
    add_bullet("លេខសម្គាល់សិស្ស (Student ID / Code)៖ ឧទាហរណ៍ STU-001 (ជាលេខកូដដែលសិស្សប្រើសម្រាប់ចូលប្រឡងអនឡាញ)។", "ជំហានទី ២ (បំពេញព័ត៌មាន)៖ ")
    add_bullet("គោត្តនាម និងនាម៖ វាយបញ្ចូលឈ្មោះជាភាសាខ្មែរ (ឧទាហរណ៍៖ កែវ វិសាល)។")
    add_bullet("ភេទ៖ ជ្រើសរើស ប្រុស ឬ ស្រី។")
    add_bullet("ថ្ងៃខែឆ្នាំកំណើត៖ ជ្រើសរើសតាមប្រតិទិន។")
    add_bullet("ថ្នាក់រៀន៖ ជ្រើសរើសថ្នាក់ដែលសិស្សនោះកំពុងរៀន។")
    add_bullet("លេខទូរស័ព្ទអាណាព្យាបាល៖ សម្រាប់ទំនាក់ទំនងករណីចាំបាច់។")
    add_bullet("ចុច «រក្សាទុកសិស្ស (Save Student)»។", "ជំហានទី ៣៖ ")

    add_heading_2("៦.២. ការច្រោះមើលឈ្មោះសិស្ស")
    add_body("លោកគ្រូ-អ្នកគ្រូអាចប្រើប្រអប់ជ្រើសរើសថ្នាក់រៀន (Filter by Class) ដើម្បីបង្ហាញតែបញ្ជីសិស្សក្នុងថ្នាក់ដែលចង់ពិនិត្យបានយ៉ាងលឿន។")

    # ================= CHAPTER 7 =================
    add_heading_1("ជំពូកទី ៧៖ ការកត់ត្រា និងតាមដានអវត្តមានសិស្ស (Attendance Tracking)")
    add_body("ការកត់ត្រាវត្តមានត្រូវបានធ្វើទំនើបកម្មឱ្យស្របតាមស្តង់ដារអប់រំនៅកម្ពុជា ដោយបែងចែកប្រភេទអវត្តមានជា «ច្បាប់» និង «គ្មានច្បាប់» យ៉ាងច្បាស់លាស់។")

    add_heading_2("៧.១. និមិត្តសញ្ញា និងប្រភេទវត្តមានថ្មី")
    add_bullet("វត្តមានពេញលេញ (Present - ពណ៌បៃតង)៖ សិស្សបានចូលរួមរៀនពេញម៉ោង។", "✔️ វត្តមាន (Present)៖ ")
    add_bullet("អវត្តមានមានច្បាប់ (Permission - សម្គាល់ដោយ «1ច» ឬ «ច»)៖ សិស្សមានលិខិតសុំច្បាប់ត្រឹមត្រូវពីអាណាព្យាបាល។", "📋 ច្បាប់ (1ច / Permission)៖ ")
    add_bullet("អវត្តមានគ្មានច្បាប់ (Unexcused - សម្គាល់ដោយ «1គ» ឬ «គ»)៖ សិស្សអវត្តមានដោយគ្មានមូលហេតុ ឬគ្មានការអនុញ្ញាត។", "❌ គ្មានច្បាប់ (1គ / Absent)៖ ")

    add_heading_2("៧.២. ជំហានកត់ត្រាវត្តមានប្រចាំថ្ងៃ")
    add_bullet("ចូលទៅកាន់ម៉ឺនុយ «បញ្ជីវត្តមាន (Attendance)»។", "ជំហានទី ១៖ ")
    add_bullet("ជ្រើសរើស ថ្នាក់រៀន និង កាលបរិច្ឆេទ (ថ្ងៃ ខែ ឆ្នាំ) ដែលត្រូវកត់ត្រា។", "ជំហានទី ២៖ ")
    add_bullet("ប្រព័ន្ធនឹងរាយនាមសិស្សទាំងអស់ក្នុងថ្នាក់នោះ។ លំនាំដើមសិស្សទាំងអស់ត្រូវបានកំណត់ជា «វត្តមាន» ជាស្រេច។", "ជំហានទី ៣៖ ")
    add_bullet("លោកគ្រូ-អ្នកគ្រូគ្រាន់តែចុចប្តូរតែលើសិស្សណាដែល «ច្បាប់ (1ច)» ឬ «គ្មានច្បាប់ (1គ)» ប៉ុណ្ណោះ។", "ជំហានទី ៤៖ ")
    add_bullet("ចុចលើប៊ូតុង «រក្សាទុកវត្តមាន (Save Attendance)»។", "ជំហានទី ៥៖ ")

    add_heading_2("៧.៣. ការបូកសរុបអវត្តមានប្រចាំខែ និងប្រចាំឆ្នាំ")
    add_body("ប្រព័ន្ធនឹងធ្វើការបូកសរុបអវត្តមានដោយស្វ័យប្រវត្តិ៖")
    add_bullet("អវត្តមានប្រចាំខែ៖ បង្ហាញចំនួនសរុបនៃ ច្បាប់ (ច) និង គ្មានច្បាប់ (គ) ក្នុងខែនីមួយៗ នៅលើតារាងពិន្ទុ និងរបាយការណ៍។")
    add_bullet("អវត្តមានសរុបប្រចាំឆ្នាំ៖ បូកសរុបវត្តមានពេញមួយឆ្នាំសិក្សា ដើម្បីវាយតម្លៃលើលក្ខខណ្ឌសិទ្ធិប្រឡង ឬការឡើងថ្នាក់។")

    # ================= CHAPTER 8 =================
    add_heading_1("ជំពូកទី ៨៖ ការបញ្ចូល និងគ្រប់គ្រងពិន្ទុសិស្ស (Scores & Grading Management)")
    add_body("ទំព័រ «បញ្ជីពិន្ទុ (Scores)» ត្រូវបានរចនាឡើងសាថ្មីជាមួយនឹងក្បាលតារាងធំ ច្បាស់ ពណ៌ដិតខ្ពស់ (High Contrast) ងាយស្រួលមើលបំផុតសម្រាប់លោកគ្រូ-អ្នកគ្រូគ្រប់វ័យ។")

    add_heading_2("៨.១. សមាសភាគពិន្ទុចម្រុះ")
    add_body("ប្រព័ន្ធគាំទ្រការវាយតម្លៃពហុសមាសភាគតាមក្បួនគរុកោសល្យ៖")
    add_bullet("តេស្តដើមឆ្នាំ (Baseline Test)៖ សម្រាប់វាស់ស្ទង់សមត្ថភាពដំបូងរបស់សិស្សនៅដើមឆ្នាំសិក្សា។", "១. ")
    add_bullet("ពិន្ទុប្រចាំខែ (Monthly Test Score)៖ ពិន្ទុតេស្ត ឬប្រឡងប្រចាំខែជាក់ស្តែង។", "២. ")
    add_bullet("ពិន្ទុវត្តមាន (Attendance Score)៖ គណនាផ្អែកលើការមកសាលាទៀងទាត់របស់សិស្ស។", "៣. ")
    add_bullet("ពិន្ទុកិច្ចការស្រាវជ្រាវ (Assignment / Research Score)៖ ពិន្ទុកិច្ចការផ្ទះ ឬការស្រាវជ្រាវស្វ័យសិក្សា។", "៤. ")

    add_heading_2("៨.២. ការកំណត់ពិន្ទុពេញ (Custom Max Score Rules)")
    add_body("លោកគ្រូ-អ្នកគ្រូអាចកំណត់ទម្ងន់ពិន្ទុពេញតាមមុខវិជ្ជា និងថ្នាក់រៀននីមួយៗបានដោយសេរី (ឧទាហរណ៍៖ តេស្ត ៧០, វត្តមាន ១៥, កិច្ចការ ១៥ ឬ តេស្ត ៥០, វត្តមាន ២០, កិច្ចការ ៣០)។ ប្រព័ន្ធនឹងគណនាផលបូក និងភាគរយដោយស្វ័យប្រវត្តិតាមការកំណត់ជាក់ស្តែងរបស់គ្រូ។")

    add_heading_2("៨.៣. តារាងកម្រិតនិទ្ទេស និងចំណាត់ថ្នាក់ស្វ័យប្រវត្តិ")
    add_body("ប្រព័ន្ធនឹងបូកពិន្ទុសរុប រួចកំណត់និទ្ទេស (Grade A-F) និងគណនា ចំណាត់ថ្នាក់លេខ ១, ២, ៣... (Rank) ជូនសិស្សដោយស្វ័យប្រវត្តិ៖")

    tbl_grade = doc.add_table(rows=1, cols=4)
    tbl_grade.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_grade.columns[0].width = Inches(1.3)
    tbl_grade.columns[1].width = Inches(1.8)
    tbl_grade.columns[2].width = Inches(1.8)
    tbl_grade.columns[3].width = Inches(1.6)

    g_hdrs = ["និទ្ទេស (Grade)", "ភាគរយពិន្ទុ (%)", "ពិន្ទុជាក់ស្តែង (លើ១០០)", "ការវាយតម្លៃ"]
    for i, h in enumerate(g_hdrs):
        c = tbl_grade.rows[0].cells[i]
        c.text = h
        set_cell_background(c, "0284C7")
        set_cell_margins(c, top=80, bottom=80, left=100, right=100)
        for r in c.paragraphs[0].runs:
            format_run(r, font_name="Khmer OS Battambang", size_pt=10, bold=True, color_rgb=(255, 255, 255))

    grades_data = [
        ("A", "៩០% - ១០០%", "៩០ - ១០០", "ល្អប្រសើរណាស់ (Excellent)"),
        ("B", "៨០% - ៨៩%", "៨០ - ៨៩", "ល្អណាស់ (Very Good)"),
        ("C", "៧០% - ៧៩%", "៧០ - ៧៩", "ល្អ (Good)"),
        ("D", "៦០% - ៦៩%", "៦០ - ៦៩", "ល្អបង្គួរ (Fair)"),
        ("E", "៥០% - ៥៩%", "៥០ - ៥៩", "មធ្យម / ជាប់ (Pass)"),
        ("F", "ក្រោម ៥០%", "ក្រោម ៥០", "ធ្លាក់ (Fail)")
    ]

    for row in grades_data:
        r_cells = tbl_grade.add_row().cells
        for i in range(4):
            r_cells[i].text = row[i]
            set_cell_background(r_cells[i], "F1F5F9" if row[0] in ["A", "B"] else ("FEF2F2" if row[0] == "F" else "FFFFFF"))
            set_cell_margins(r_cells[i], top=60, bottom=60, left=100, right=100)
            for r in r_cells[i].paragraphs[0].runs:
                format_run(r, font_name="Khmer OS Battambang", size_pt=10, bold=(i == 0), color_rgb=(239, 68, 68) if row[0] == "F" else (30, 41, 59))

    # ================= CHAPTER 9 =================
    add_heading_1("ជំពូកទី ៩៖ ស្ថិតិ និងការវិភាគទិន្នន័យសិក្សា (Student Statistics)")
    add_body("ទំព័រ «ស្ថិតិសិស្ស (Statistics)» ជួយឱ្យលោកគ្រូ-អ្នកគ្រូ និងគណៈគ្រប់គ្រងអាចវាយតម្លៃគុណភាពនៃការបង្រៀន និងការរៀនបានយ៉ាងស៊ីជម្រៅ តាមរយៈតារាងទិន្នន័យ និងភាគរយជាក់ស្តែង។")

    add_heading_2("៩.១. តារាងស្ថិតិតាមមុខវិជ្ជានីមួយៗប្រចាំខែ (Subject-by-Subject Monthly Breakdown)")
    add_body("នៅក្នុងការកែលម្អសាថ្មី ប្រព័ន្ធបានបន្ថែម «តារាងស្ថិតិតាមមុខវិជ្ជានីមួយៗប្រចាំខែ» ដែលមានក្បាលតារាងធំ ច្បាស់ និងទិន្នន័យលម្អិត៖")
    add_bullet("មុខវិជ្ជា និងខែសិក្សា។")
    add_bullet("ចំនួនសិស្សសរុប និងចំនួនសិស្សស្រី។")
    add_bullet("មធ្យមភាគពិន្ទុរួមរបស់ថ្នាក់។")
    add_bullet("ចំនួនសិស្សទទួលបាននិទ្ទេស A, B, C, D, E, F (រាប់ទាំងសរុប និងសិស្សស្រី)។")
    add_bullet("ចំនួនសិស្សជាប់សរុប សិស្សស្រីជាប់ និងភាគរយសិស្សជាប់ (%)។")

    add_heading_2("៩.២. ស្ថិតិអត្រាសិស្សជាប់ និងធ្លាក់")
    add_body("បង្ហាញជាភាគរយជាក់ស្តែងនូវអត្រាសិស្សដែលទទួលបានពិន្ទុចាប់ពី ៥០% ឡើងទៅ (ជាប់) និងសិស្សក្រោម ៥០% (ធ្លាក់) ដើម្បីងាយស្រួលក្នុងការរៀបចំផែនការបំប៉នសិស្សខ្សោយទាន់ពេលវេលា។")

    # ================= CHAPTER 10 =================
    add_heading_1("ជំពូកទី ១០៖ វិញ្ញាបនបត្របញ្ជាក់ការសិក្សា និងការផ្ទៀងផ្ទាត់ QR Code")
    add_body("ម៉ូឌុល «វិញ្ញាបនបត្រ (Certificates)» គឺជាសមិទ្ធផលដ៏លេចធ្លោនៃប្រព័ន្ធគ្រប់គ្រងថ្នាក់រៀន ២០២៧ ដែលអនុញ្ញាតឱ្យសាលារៀន និងគ្រូបង្រៀនអាចចេញវិញ្ញាបនបត្រឌីជីថលផ្លូវការប្រកបដោយស្តង់ដារអន្តរជាតិ។")

    add_heading_2("១០.១. របៀបចេញវិញ្ញាបនបត្រជូនសិស្ស")
    add_bullet("ចូលទៅកាន់ម៉ឺនុយ «វិញ្ញាបនបត្រ» នៅលើ Sidebar ខាងឆ្វេង។", "ជំហានទី ១៖ ")
    add_bullet("ជ្រើសរើស ថ្នាក់រៀន និងជ្រើសរើស ឈ្មោះសិស្ស ដែលត្រូវទទួលវិញ្ញាបនបត្រ។", "ជំហានទី ២៖ ")
    add_bullet("ប្រព័ន្ធនឹងរៀបចំទម្រង់វិញ្ញាបនបត្រយ៉ាងស្រស់ស្អាតដោយស្វ័យប្រវត្តិ។", "ជំហានទី ៣៖ ")
    add_bullet("នៅលើវិញ្ញាបនបត្រមានភ្ជាប់មកជាមួយ៖ ឈ្មោះសិស្ស, ថ្នាក់រៀន, សាលារៀន, ឆ្នាំសិក្សា, ត្រាសាលា, ហត្ថលេខានាយក, ហត្ថលេខាគ្រូ, និងលេខកូដស៊េរីវិញ្ញាបនបត្រ (Certificate Serial Number)។", "ជំហានទី ៤៖ ")
    add_bullet("ចុចប៊ូតុង «បោះពុម្ពវិញ្ញាបនបត្រ (Print Certificate)» ដើម្បីព្រីនជាក្រដាស A4 ឬរក្សាទុកជា PDF។", "ជំហានទី ៥៖ ")

    add_heading_2("១០.២. ប្រព័ន្ធ QR Code ផ្ទៀងផ្ទាត់សច្ចភាពអនឡាញ (Instant Online Verification)")
    add_body("វិញ្ញាបនបត្រនីមួយៗមានបង្កប់នូវ QR Code ពិសេស៖")
    add_bullet("អាណាព្យាបាល សិស្ស ឬស្ថាប័នពាក់ព័ន្ធ អាចយកទូរស័ព្ទដៃស្កេនលើ QR Code នោះ។")
    add_bullet("ប្រព័ន្ធនឹងបើកទំព័រផ្លូវការ៖ https://class.mitpukae.com/certificate/verify/{serial} ភ្លាមៗ។")
    add_bullet("ទំព័រផ្ទៀងផ្ទាត់នឹងបង្ហាញព័ត៌មានលម្អិតពិតប្រាកដរបស់សិស្ស ឈ្មោះសាលា ឆ្នាំសិក្សា ព្រមទាំងរូបភាពត្រា និងហត្ថលេខាដើម ដែលការពារការក្លែងបន្លំ ឬកែសម្រួលឯកសារខុសច្បាប់បាន ១០០%!")

    # ================= CHAPTER 11 =================
    add_heading_1("ជំពូកទី ១១៖ ការគ្រប់គ្រងវិញ្ញាសា និងការប្រឡងអនឡាញ (Online Exams)")
    add_body("ប្រព័ន្ធបានបំពាក់ម៉ូឌុលប្រឡងអនឡាញ ដែលអនុញ្ញាតឱ្យលោកគ្រូ-អ្នកគ្រូអាចបង្កើតវិញ្ញាសា និងឱ្យសិស្សធ្វើតេស្តតាមទូរស័ព្ទ ឬកុំព្យូទ័របានដោយស្វ័យប្រវត្តិ។")

    add_heading_2("១១.១. របៀបដែលគ្រូបង្កើតវិញ្ញាសា")
    add_bullet("ចូលទៅកាន់ម៉ឺនុយ «វិញ្ញាសាប្រឡង (Exams)»។", "ជំហានទី ១៖ ")
    add_bullet("ជ្រើសរើស ខែប្រឡង និង ថ្នាក់រៀន។", "ជំហានទី ២៖ ")
    add_bullet("បញ្ចូលសំណួរ និងជម្រើសចម្លើយ (A, B, C, D) រួមទាំងកំណត់ចម្លើយដែលត្រឹមត្រូវ។", "ជំហានទី ៣៖ ")
    add_bullet("ចុច «រក្សាទុកវិញ្ញាសា»។ វិញ្ញាសានឹងបើកឱ្យសិស្សក្នុងថ្នាក់នោះចូលប្រឡងបានភ្លាមៗ។", "ជំហានទី ៤៖ ")

    add_heading_2("១១.២. របៀបដែលសិស្សចូលប្រឡង (Student Exam Portal)")
    add_bullet("សិស្សបើកទូរស័ព្ទដៃ ឬកុំព្យូទ័រ ហើយចូលទៅកាន់ Link៖ https://class.mitpukae.com/test", "ជំហានទី ១៖ ")
    add_bullet("សិស្សបញ្ចូល លេខកូដសិស្ស របស់ខ្លួន (ឧទាហរណ៍៖ STU-001) រួចចុច «ចូលប្រឡង»។", "ជំហានទី ២៖ ")
    add_bullet("សិស្សអានសំណួរ និងចុចជ្រើសរើសចម្លើយរហូតដល់ចប់ រួចចុច «បញ្ជូនចម្លើយ (Submit)»។", "ជំហានទី ៣៖ ")
    add_bullet("ប្រព័ន្ធនឹងកែពិន្ទុ និងបង្ហាញលទ្ធផលជូនសិស្សភ្លាមៗ ហើយបញ្ជូនទិន្នន័យពិន្ទុនេះចូលទៅកាន់បញ្ជីពិន្ទុរបស់គ្រូដោយស្វ័យប្រវត្តិ។", "ជំហានទី ៤៖ ")

    # ================= CHAPTER 12 =================
    add_heading_1("ជំពូកទី ១២៖ ការគ្រប់គ្រងមេរៀន និងកាលវិភាគបង្រៀន (Lessons & Schedules)")
    add_body("សម្រួលដល់ការរៀបចំឯកសារបង្រៀន និងកាលវិភាគប្រចាំសប្តាហ៍ឱ្យមានរបៀបរៀបរយ។")

    add_heading_2("១២.១. ការគ្រប់គ្រង និងចែករំលែកមេរៀន (Lesson Management)")
    add_body("លោកគ្រូ-អ្នកគ្រូអាចបង្ហោះមេរៀន ស្លាយបង្រៀន (PowerPoint), ឯកសារ PDF ឬតំណភ្ជាប់វីដេអូបង្រៀន ដើម្បីឱ្យសិស្សអាចចូលមើល និងទាញយកទៅស្វ័យសិក្សានៅផ្ទះបាន។")
    add_bullet("ការបង្ហោះមេរៀន៖ ចុច «មើលមេរៀន» -> បញ្ចូលចំណងជើងមេរៀន មាតិកា និង Upload ឯកសារភ្ជាប់។")
    add_bullet("ការកំណត់សិទ្ធិ៖ អាចជ្រើសរើសចែករំលែកសម្រាប់ «គ្រប់ថ្នាក់» ឬសម្រាប់តែ «ថ្នាក់ជាក់លាក់ណាមួយ»។")

    add_heading_2("១២.២. ការរៀបចំកាលវិភាគបង្រៀន (Teaching Schedule)")
    add_body("កាលវិភាគត្រូវបានរៀបចំឡើងជាតារាងប្រចាំសប្តាហ៍ (ពីថ្ងៃចន្ទ ដល់ ថ្ងៃសៅរ៍/អាទិត្យ)៖")
    add_bullet("បង្ហាញម៉ោងបង្រៀនច្បាស់លាស់ (ម៉ោងពេលព្រឹក ពេលរសៀល ឬពេលល្ងាច)។")
    add_bullet("បង្ហាញមុខវិជ្ជា និងបន្ទប់សិក្សានីមួយៗ ដើម្បីកុំឱ្យមានការជាន់ម៉ោងគ្នារវាងគ្រូនិងគ្រូ។")

    # ================= CHAPTER 13 =================
    add_heading_1("ជំពូកទី ១៣៖ ការទាញយក និងបោះពុម្ពរបាយការណ៍សរុប (Reports & Printing)")
    add_body("នៅដំណាច់ខែនីមួយៗ លោកគ្រូ-អ្នកគ្រូត្រូវចេញរបាយការណ៍សរុបជូនគណៈគ្រប់គ្រង និងអាណាព្យាបាលសិស្ស។ ទំព័រ «របាយការណ៍ (Report)» ត្រូវបានរចនាឡើងយ៉ាងពិសេសសម្រាប់ការងារនេះ។")

    add_heading_2("១៣.១. មាតិកាក្នុងតារាងរបាយការណ៍សរុប")
    add_bullet("ចំណាត់ថ្នាក់សិស្ស (Rank)៖ ចំណាត់ថ្នាក់លេខ ១, ២, ៣... គណនាដោយស្វ័យប្រវត្តិតាមពិន្ទុសរុបខ្ពស់ទៅទាប។")
    add_bullet("ព័ត៌មានសិស្ស៖ លេខសម្គាល់, គោត្តនាម-នាម, ភេទ។")
    add_bullet("លទ្ធផលសិក្សា៖ ពិន្ទុសរុប និងនិទ្ទេសរួម (Grade A-F)។")
    add_bullet("កំណត់ត្រាវត្តមាន៖ ចំនួនថ្ងៃមានវត្តមាន និងចំនួនអវត្តមាន (ច្បាប់ និងគ្មានច្បាប់) ក្នុងខែនោះ។")

    add_heading_2("១៣.២. ការបោះពុម្ពរបាយការណ៍ជាក្រដាស A4 ឬ PDF")
    add_bullet("ជ្រើសរើស ថ្នាក់រៀន និង ខែ ដែលចង់ចេញរបាយការណ៍។", "ជំហានទី ១៖ ")
    add_bullet("ពិនិត្យមើលតារាងទិន្នន័យឱ្យបានរួចរាល់។", "ជំហានទី ២៖ ")
    add_bullet("ចុចលើប៊ូតុងពណ៌ខៀវ «បោះពុម្ពរបាយការណ៍ (Print 🖨️)» នៅផ្នែកខាងលើ។", "ជំហានទី ៣៖ ")
    add_bullet("ផ្ទាំង Print របស់កុំព្យូទ័រនឹងលោតឡើង៖", "ជំហានទី ៤៖ ")
    add_bullet("ប្រសិនបើចង់បោះពុម្ពជាក្រដាស៖ ជ្រើសរើសម៉ាស៊ីនព្រីន (Printer) របស់អ្នក រួចចុច Print។")
    add_bullet("ប្រសិនបើចង់រក្សាទុកជាឯកសារ PDF៖ នៅកន្លែង Destination សូមជ្រើសរើស «Save as PDF» រួចចុច Save។")

    add_callout("ភាពស្អាតនៃទម្រង់បោះពុម្ព (Clean Print Layout)", "ប្រព័ន្ធបានកំណត់កូដពិសេស (.no-print)។ នៅពេលលោកគ្រូ-អ្នកគ្រូចុចបោះពុម្ព គឺមានតែតារាងរបាយការណ៍ផ្លូវការប៉ុណ្ណោះដែលបង្ហាញនៅលើក្រដាស A4 ចំណែកឯម៉ឺនុយ Sidebar, ប៊ូតុងចុចផ្សេងៗ និងផ្ទាំង Telegram នឹងត្រូវលាក់ស្វ័យប្រវត្តិ ធានាបាននូវទម្រង់របាយការណ៍ផ្លូវការដ៏ស្រស់ស្អាត ១០០%!")

    # ================= CHAPTER 14 =================
    add_heading_1("ជំពូកទី ១៤៖ ការគ្រប់គ្រងគណនីគ្រូបង្រៀនសម្រាប់ Admin (Admin Accounts Management)")
    add_body("ផ្ទាំង «គ្រប់គ្រងគណនីគ្រូបង្រៀន (/admin/accounts)» ផ្តល់ជូនគណៈគ្រប់គ្រងសាលា និង Admin នូវលទ្ធភាពគ្រប់គ្រងបុគ្គលិកបង្រៀនទាំងអស់ក្នុងប្រព័ន្ធយ៉ាងពេញលេញ និងមានសុវត្ថិភាព។")

    add_heading_2("១៤.១. ការចូលកាន់ផ្ទាំង Admin")
    add_body("មានតែគណនីដែលមានសិទ្ធិ ROLE_ADMIN ប៉ុណ្ណោះដែលអាចចូលប្រើប្រាស់ផ្ទាំងនេះបានតាមរយៈម៉ឺនុយ «គ្រប់គ្រងគណនី» ឬអាសយដ្ឋាន៖ https://class.mitpukae.com/admin/accounts")

    add_heading_2("១៤.២. តារាងព័ត៌មានលម្អិតរបស់គ្រូបង្រៀន")
    add_body("តារាងគ្រប់គ្រងគណនីត្រូវបានកែលម្អយ៉ាងច្បាស់លាស់ ដោយបង្ហាញព័ត៌មានសំខាន់ៗ៖")
    add_bullet("ឈ្មោះពេញ និងតួនាទី (FullName & Role)។")
    add_bullet("ឈ្មោះអ្នកប្រើប្រាស់ (Username)។")
    add_bullet("មុខវិជ្ជាបង្រៀន (Teaching Subject)។")
    add_bullet("ថ្នាក់បង្រៀន (Teaching Classes)។")
    add_bullet("លេខទូរស័ព្ទ (Phone Number)។")
    add_bullet("គណនី Telegram (Telegram Username/Link)។")
    add_bullet("ស្ថានភាពគណនី (Status: Active ដំណើរការ ឬ Suspended ផ្អាក)។")

    add_heading_2("១៤.៣. ការបង្កើតគណនីគ្រូថ្មី និងការកំណត់លេខសម្ងាត់")
    add_bullet("ចុចលើប៊ូតុង «+ បង្កើតគណនីថ្មី»។", "ជំហានទី ១៖ ")
    add_bullet("បំពេញព័ត៌មានគ្រូ៖ Username, Password, FullName, Subject, Classes, Phone, Telegram។", "ជំហានទី ២៖ ")
    add_bullet("ចុច «រក្សាទុក»។ គ្រូបង្រៀនអាចយក Username និង Password នោះទៅ Login បានភ្លាមៗ។", "ជំហានទី ៣៖ ")
    add_bullet("ការប្តូរលេខសម្ងាត់ (Reset Password)៖ Admin អាចចុចលើប៊ូតុង «ប្តូរលេខសម្ងាត់» នៅជួរគ្រូនីមួយៗ ដើម្បីផ្តល់លេខសម្ងាត់ថ្មីជូនគ្រូករណីភ្លេច។", "ជំហានទី ៤៖ ")

    add_heading_2("១៤.៤. ការបើក ឬផ្អាកដំណើរការគណនី (Toggle Status)")
    add_body("Admin អាចចុចប៊ូតុងបិទ/បើកសិទ្ធិ (Active / Suspended) ដោយគ្រាន់តែចុចមួយឃ្លីកប៉ុណ្ណោះ។ ប្រសិនបើគណនីត្រូវបានផ្អាក គ្រូនោះនឹងមិនអាច Login ចូលប្រព័ន្ធបានឡើយ។")

    # ================= CHAPTER 15 =================
    add_heading_1("ជំពូកទី ១៥៖ ផ្នែកទំនាក់ទំនង ជំនួយបច្ចេកទេស និងសុវត្ថិភាព")
    add_body("ដើម្បីធានាការប្រើប្រាស់ប្រកបដោយទំនុកចិត្ត និងដោះស្រាយរាល់ចម្ងល់បច្ចេកទេសបានទាន់ពេលវេលា ប្រព័ន្ធបានភ្ជាប់ទំនាក់ទំនងផ្ទាល់ជាមួយក្រុមការងារបច្ចេកទេស។")

    add_heading_2("១៥.១. ព័ត៌មានទំនាក់ទំនងជំនួយផ្ទាល់")
    add_bullet("លោកគ្រូ សំ គឹមឆេង (Sam Kimchheng)", "👤 អ្នកគ្រប់គ្រងបច្ចេកទេស៖ ")
    add_bullet("https://t.me/Kimchheng12", "✈️ តំណភ្ជាប់ Telegram ផ្ទាល់៖ ")
    add_bullet("@Kimchheng12", "💬 ឈ្មោះគណនី Telegram (Username)៖ ")

    add_heading_2("១៥.២. គន្លឹះសុវត្ថិភាពទិន្នន័យ")
    add_bullet("កុំចែករំលែកលេខសម្ងាត់គណនីគ្រូរបស់អ្នកទៅកាន់អ្នកដទៃ ឬសិស្ស។")
    add_bullet("ប្រសិនបើលោកគ្រូ-អ្នកគ្រូប្រើប្រាស់កុំព្យូទ័ររួម (Computer Lab សាលា ឬហាងអ៊ីនធឺណិត) សូមចុច «ចាកចេញ (Logout 🚪)» រាល់ពេលប្រើប្រាស់រួច។")

    add_heading_2("១៥.៣. របៀបដោះស្រាយនៅពេលប្រព័ន្ធមិនទាន់បង្ហាញទិន្នន័យថ្មី")
    add_bullet("ចុចបញ្ជា «Ctrl + F5» លើកុំព្យូទ័រ (ឬចុច Reload លើទូរស័ព្ទដៃ) ដើម្បីឱ្យ Browser ទាញយកទិន្នន័យថ្មីស្រឡាងពី Server។")
    add_bullet("ប្រសិនបើមានបញ្ហាបច្ចេកទេសណាមួយ សូម Screenshot អេក្រង់រួចផ្ញើមកកាន់ Telegram @Kimchheng12 ដើម្បីទទួលបានការជួយសម្រួលភ្លាមៗ។")

    # Save to all destination paths
    paths = [
        r"D:\NUBB DOCUMENTARY\Class Management System2027\សៀវភៅណែនាំ_ការប្រើប្រាស់ប្រព័ន្ធគ្រប់គ្រងថ្នាក់រៀន_២០២៧.docx",
        r"D:\NUBB DOCUMENTARY\Class Management System2027\សៀវភៅណែនាំ_ការប្រើប្រាស់ប្រព័ន្ធគ្រប់គ្រងថ្នាក់រៀន_២០២៧.doc",
        r"D:\NUBB DOCUMENTARY\សៀវភៅណែនាំ_ការប្រើប្រាស់ប្រព័ន្ធគ្រប់គ្រងថ្នាក់រៀន_២០២៧.docx",
        r"D:\NUBB DOCUMENTARY\សៀវភៅណែនាំ_ការប្រើប្រាស់ប្រព័ន្ធគ្រប់គ្រងថ្នាក់រៀន_២០២៧.doc",
        r"D:\NUBB DOCUMENTARY\Class Management System2027\Class_Management_System_Teacher_Guide_2027.docx",
        r"D:\NUBB DOCUMENTARY\Class_Management_System_Teacher_Guide_2027.docx",
        r"D:\NUBB DOCUMENTARY\Class Management System2027\src\main\resources\static\docs\teacher-guide.docx",
        r"D:\NUBB DOCUMENTARY\Class Management System2027\src\main\resources\static\docs\teacher-guide.doc",
    ]
    
    os.makedirs(r"D:\NUBB DOCUMENTARY\Class Management System2027\src\main\resources\static\docs", exist_ok=True)

    for p in paths:
        try:
            doc.save(p)
            print("Successfully saved to:", p.encode('ascii', 'replace').decode('ascii'))
        except Exception as e:
            print("Error saving to", p.encode('ascii', 'replace').decode('ascii'), ":", e)

if __name__ == "__main__":
    create_document()
